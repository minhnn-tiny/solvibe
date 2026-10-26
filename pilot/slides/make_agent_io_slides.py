"""Build the deep-dive deck on what each AI model got, and gave back, in the finding 63008 example.

The control group deck is the base, so the theme, layouts and fonts match. Its slides are removed
and new ones are built: one slide per AI step, each followed by slides with the actual data (prompts,
requirements, agent diffs, model replies and check tests), read from the run files at build time.
The facts come from the files named in each slide's source line:
  pilot/tasks/63008/feature_tasks/results.md  (agent results, review, repair round and judge)
  pilot/tasks/63008/control_group.md          (requirements and PoVs)
  pilot/prompts/*.md and pilot/scripts/run_task.py (what each model is sent)

Usage (python-pptx is not in the project environment):
  uv run --with python-pptx python pilot/slides/make_agent_io_slides.py \
      "slides/Finding 63008 Control Group Slides.pptx" "slides/Finding 63008 AI Agent Inputs and Outputs.pptx"
"""
import json
import re
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn, nsdecls

SRC, DST = sys.argv[1], sys.argv[2]
# The control group deck's colours and fonts.
INK, MUTED, BLUE, RED, GREEN, AMBER = "18202B", "566173", "2C4F86", "B42318", "1D7A46", "8A5A00"
BG, CARD, LINE, HEAD, GRID, QUOTE = "F6F7F9", "FCFCFD", "D9DEE6", "E9E9EA", "C9C9CA", "ECEFF3"
RED_BG, GREEN_BG, BLUE_BG, AMBER_BG, DARK = "FDECEA", "E7F5EC", "E6EAF0", "FFF4DC", "1B2638"
ON_DARK, ON_DARK_MUTED, ON_DARK_KICKER = "F6F7F9", "C8D1DE", "9DB8E4"
# Code is in Courier New, not the control group deck's JetBrains Mono: JetBrains Mono joins != and <= into
# single symbols, which would misquote the code on the data slides.
HEADF, BODYF, MONO = "Archivo", "Source Sans 3", "Courier New"
L, W = 0.89, 11.55
# The three columns of a step slide: what the model gets, the model, what it gives back.
XI, WI, XA, WA, XO, WO = 0.89, 4.45, 5.68, 1.9, 7.92, 4.52
CY, CH = 1.92, 3.76

prs = Presentation(SRC)
layout = prs.slides[0].slide_layout
for sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)


def rgb(h): return RGBColor.from_string(h)


def set_font(run, face, size, color, bold=None, italic=None, spacing=None):
    f = run.font
    f.name, f.size, f.bold, f.italic = face, Pt(size), bold, italic
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    if spacing:
        rpr.set("spc", str(spacing))
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", face)


def fill_frame(tf, paras, face, size, color, bold, align, anchor, bullets, after, line):
    """paras: a str, or a list of (text, options) runs, per paragraph."""
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(line or size * 1.32)
        if after and i < len(paras) - 1:
            p.space_after = Pt(after)
        if bullets:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "155194"); ppr.set("indent", "-155194")
            ppr.append(ppr.makeelement(qn("a:buSzPct"), {"val": "100000"}))
            ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            set_font(r, o.get("face", face), o.get("size", size), o.get("color", color),
                     o.get("bold", bold), o.get("italic"))


def text(slide, x, y, w, h, paras, face=BODYF, size=12, color=INK, bold=None, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, after=0, line=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_frame(tb.text_frame, paras, face, size, color, bold, align, anchor, bullets, after, line)
    return tb


def shape(slide, kind, x, y, w, h, fill, line=None, radius=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        s.adjustments[0] = radius
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(0.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, fill=CARD, line=LINE):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius=0.055)


def pill(slide, x, y, w, label, fill, color):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.26, fill, radius=0.5)
    fill_frame(s.text_frame, [label], HEADF, 11, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, False, 0, None)
    return s


def quote(slide, x, y, w, h, paras, face=BODYF, size=11):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, QUOTE, radius=0.06)
    text(slide, x + 0.12, y + 0.08, w - 0.24, h - 0.16, paras, face=face, size=size, color=INK)


def arrow(slide, x, y, w=0.22, h=0.32):
    shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, MUTED)


def label(slide, x, y, w, words, color=BLUE):
    text(slide, x, y, w, 0.26, [[(words, {})]], face=HEADF, size=12, color=color, bold=True)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("spc", "100")


def table(slide, x, y, widths, rows, row_h=0.4, size=12, colors=None, faces=None, pad=0.06):
    """rows[0] is the header. colors and faces map (row, col) to a text colour or font. A cell can also
    hold a list of paragraphs, each a list of (text, options) runs, which then sets its own fonts."""
    colors, faces = colors or {}, faces or {}
    gf = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                                Inches(sum(widths)), Inches(row_h * len(rows)))
    tbl = gf.table
    tbl.first_row = tbl.horz_banding = False
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(pad)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid(); c.fill.fore_color.rgb = rgb(HEAD if i == 0 else CARD)
            tf = c.text_frame
            tf.word_wrap = True
            if isinstance(val, list):
                fill_frame(tf, val, BODYF, size, MUTED, None, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, False, 2, size * 1.22)
            else:
                p = tf.paragraphs[0]
                p.line_spacing = Pt(size * 1.3)
                r = p.add_run()
                r.text = val
                set_font(r, faces.get((i, j), BODYF), size, colors.get((i, j), INK if i == 0 or j == 0 else MUTED),
                         bold=(i == 0 or j == 0) or None)
            tcpr = c._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcpr.makeelement(qn(side), {"w": "9525"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": GRID}))
                ln.append(sf)
                tcpr.append(ln)
            # Borders must come before the fill in tcPr.
            fill_el = tcpr.find(qn("a:solidFill"))
            tcpr.remove(fill_el); tcpr.append(fill_el)
    return gf


# The deck's notes master has no body placeholder, so python-pptx cannot make one. This is the
# placeholder PowerPoint wrote into the deck's own notes slides.
NOTES_BODY = ('<p:sp %s><p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/><p:cNvSpPr><a:spLocks noGrp="1"/>'
              '</p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/>'
              '<a:lstStyle/><a:p/></p:txBody></p:sp>' % nsdecls("p", "a"))


def add_notes(slide, words):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(parse_xml(NOTES_BODY))
    ns.notes_text_frame.text = words


page = [1]


def new_slide(kicker, title, source, notes):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != 1:  # keep only the title placeholder
            ph._element.getparent().remove(ph._element)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    label(s, L, 0.89, W, kicker)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(L), Inches(1.21), Inches(W), Inches(0.46)
    fill_frame(t.text_frame, [title], HEADF, 28, INK, True, PP_ALIGN.LEFT, MSO_ANCHOR.TOP, False, 0, 30.8)
    text(s, L, 6.82, 10.9, 0.26, [source], color=MUTED)
    page[0] += 1
    text(s, 11.94, 6.82, 0.5, 0.26, [str(page[0])], color=MUTED, align=PP_ALIGN.RIGHT)
    add_notes(s, notes)
    return s


def three_columns(s, get_title, model_title, give_title="OUTPUT"):
    """The frame of a step slide. Returns nothing; each slide fills the three cards itself."""
    card(s, XI, CY, WI, CH)
    label(s, XI + 0.2, CY + 0.18, WI - 0.4, get_title)
    arrow(s, XI + WI + 0.06, CY + CH / 2 - 0.16)
    card(s, XA, CY, WA, CH, fill=DARK, line=DARK)
    label(s, XA + 0.18, CY + 0.18, WA - 0.36, model_title, color=ON_DARK_KICKER)
    arrow(s, XA + WA + 0.06, CY + CH / 2 - 0.16)
    card(s, XO, CY, WO, CH)
    label(s, XO + 0.2, CY + 0.18, WO - 0.4, give_title)


def model_text(s, paras, y=None):
    text(s, XA + 0.18, y or CY + 0.56, WA - 0.36, CH - 0.7, paras, color=ON_DARK, after=7)


def band(s, heading, body, sub=None, fill=CARD, line=LINE, heading_color=BLUE, y=5.84, h=0.8):
    """A full-width card under the columns: the heading on the left, up to three lines of text on the right."""
    card(s, L, y, W, h, fill=fill, line=line)
    label(s, L + 0.2, y + 0.16, 2.3, heading, color=heading_color)
    if sub:
        text(s, L + 0.2, y + 0.44, 2.3, 0.26, [sub], color=MUTED)
    text(s, L + 2.7, y + 0.07, W - 2.9, h - 0.14, [body], anchor=MSO_ANCHOR.MIDDLE)


# ---------- helpers for the slides with the actual data ----------
MONO_EM = 0.6  # Courier New advance width, in em
BOTTOM = 6.68  # content must end above the source line
# A {{placeholder}} or an XML-style <tag> in a prompt, shown in blue monospace.
TOKEN = re.compile(r"(\{\{.*?\}\}|</?[a-z_]+(?: [a-z_]+=\"[^\"]*\")*>|\[\.\.\..*?\.\.\.\])")


def read(path):
    with open(path) as f:
        return f.read()


def fits(bottom, what):
    if bottom > BOTTOM:
        raise ValueError(f"{what} ends at {bottom:.2f} in, below {BOTTOM} in")


def diff_lines(path, ctx=1, dedent=4):
    """The lines of a patch after its first hunk header, as (kind, text) pairs. Context is cut to ctx
    lines around each change, each cut is marked with a "gap" line, and the library indent is removed."""
    body = read(path).rstrip("\n").split("\n")
    start = next(i for i, l in enumerate(body) if l.startswith("@@")) + 1
    items = []
    for l in body[start:]:
        if l.startswith("@@"):
            items.append(("gap", "..."))
            continue
        kind, t = (l[0], l[1:]) if l else (" ", "")
        items.append((kind, t[dedent:] if t.startswith(" " * dedent) else t))
    changed = [i for i, (k, _) in enumerate(items) if k in "+-"]
    out, prev = [], None
    for i in range(len(items)):
        if not any(abs(i - c) <= ctx for c in changed):
            continue
        if prev is not None and i > prev + 1:
            out.append(("gap", "..."))
        out.append(items[i])
        prev = i
    return out


def file_lines(path, first, last):
    """Lines first to last (counted from 1) of a file, unchanged."""
    return [(" ", l) for l in read(path).split("\n")[first - 1:last]]


def code_width(lines, size=9, pad=0.14):
    """The narrowest panel width that fits every line without wrapping."""
    return (max(len(t) for _, t in lines) + 2) * MONO_EM * size / 72 + 2 * pad + 0.02


def code_panel(s, x, y, w, lines, title=None, size=9, pad=0.14, min_h=0):
    """A card of code. Added lines are tinted green and removed lines red, so no line may wrap:
    raise if one is too long for the width. Returns the bottom edge of the card."""
    lh = size * 1.2 / 72
    fit = int((w - 2 * pad) / (MONO_EM * size / 72))
    for _, t in lines:
        if len(t) + 2 > fit:
            raise ValueError(f"code line of {len(t) + 2} chars, the panel fits {fit}: {t}")
    top = y + (0.38 if title else 0) + 0.08
    h = max(top - y + len(lines) * lh + 0.1, min_h)
    card(s, x, y, w, h)
    if title:
        text(s, x + pad, y + 0.13, w - 2 * pad, 0.26, [title], size=11, bold=True)
    for i, (k, _) in enumerate(lines):
        if k in "+-":
            shape(s, MSO_SHAPE.RECTANGLE, x + 0.02, top + i * lh, w - 0.04, lh, GREEN_BG if k == "+" else RED_BG)
        elif k == "h":  # a line the slide points at
            shape(s, MSO_SHAPE.RECTANGLE, x + 0.02, top + i * lh, w - 0.04, lh, AMBER_BG)
    colors = {"+": GREEN, "-": RED, " ": INK, "gap": MUTED, "c": MUTED, "h": INK}
    paras = [[((k + " " if k in "+-" else "  ") + t, {"color": colors[k]})] for k, t in lines]
    text(s, x + pad, top, w - 2 * pad, len(lines) * lh, paras, face=MONO, size=size, line=size * 1.2)
    return y + h


def prompt_runs(line, size):
    """One line of a prompt, with its placeholders and tags in blue monospace."""
    return [(p, {"face": MONO, "size": size - 1, "color": BLUE}) if TOKEN.fullmatch(p) else (p, {})
            for p in TOKEN.split(line) if p]


def verbatim(s, x, y, w, h, lines, size=9, gap=4):
    """Text copied as it is, on a tinted card. A blank line becomes a small gap, not a full line."""
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, QUOTE, radius=0.03)
    paras, gaps = [], set()
    for l in lines:
        if l.strip():
            paras.append(prompt_runs(l, size))
        elif paras:
            gaps.add(len(paras) - 1)
    tb = text(s, x + 0.14, y + 0.12, w - 0.28, h - 0.24, paras, size=size, line=size * 1.22)
    for i in gaps:
        tb.text_frame.paragraphs[i].space_after = Pt(gap)


def prompt_text(path, heading):
    """The lines of the first fenced block after a heading in a prompt file. The block ends at a line
    that is its own fence, so a prompt fenced with four backticks can hold ``` blocks."""
    lines = read(path).split("\n")
    i = next(k for k, l in enumerate(lines) if l.startswith(heading))
    i = next(k for k in range(i, len(lines)) if lines[k].startswith("```"))
    fence = lines[i][:len(lines[i]) - len(lines[i].lstrip("`"))]
    return lines[i + 1:next(k for k in range(i + 1, len(lines)) if lines[k] == fence)]


def split_columns(lines, n=3):
    """Split lines into n columns of about the same length, only at blank lines."""
    total, cols, cur, done = sum(len(l) + 40 for l in lines), [], [], 0
    for l in lines:
        cur.append(l)
        done += len(l) + 40
        if not l.strip() and len(cols) < n - 1 and done >= total * (len(cols) + 1) / n:
            cols.append(cur[:-1])
            cur = []
    return cols + [cur]


def prompt_columns(s, lines, y=1.86, size=9, starts=None, last_h=None):
    """A prompt in three columns. starts: the first words of columns 2 and 3, when splitting by length
    does not balance them. last_h: a shorter third column, to leave room under it."""
    if starts:
        cuts = [next(i for i, l in enumerate(lines) if l.startswith(w)) for w in starts]
        cols = [lines[a:b] for a, b in zip([0] + cuts, cuts + [len(lines)])]
    else:
        cols = split_columns(lines)
    gap = 0.15
    cw = (W - 2 * gap) / 3
    for i, c in enumerate(cols):
        verbatim(s, L + i * (cw + gap), y, cw, last_h if i == 2 and last_h else BOTTOM - y, c, size=size)
    return L + 2 * (cw + gap), cw


def md_runs(line, size, **extra):
    """One Markdown line, with `code` spans in monospace."""
    return [(p[1:-1], dict(extra, face=MONO, size=size - 1)) if p.startswith("`") and len(p) > 2 else (p, dict(extra))
            for p in re.split(r"(`[^`]+`)", line) if p]


def requirement(s, x, y, w, h, md, size=10.5):
    """A requirement file as the agent got it: the Markdown title in bold, code in monospace."""
    card(s, x, y, w, h)
    paras, gaps, in_code = [], set(), False
    for l in md.rstrip("\n").split("\n"):
        if l.startswith("```"):
            in_code = not in_code
        elif not l.strip():
            if paras:
                gaps.add(len(paras) - 1)
        elif in_code:
            paras.append([(l, {"face": MONO, "size": size - 1, "color": BLUE})])
        elif l.startswith("# "):
            paras.append(md_runs(l[2:], size + 1.5, bold=True))
        else:
            paras.append(md_runs(l, size))
    tb = text(s, x + 0.2, y + 0.16, w - 0.4, h - 0.3, paras, size=size, line=size * 1.28)
    for i in gaps:
        tb.text_frame.paragraphs[i].space_after = Pt(size * 0.55)


def reply(s, x, y, w, h, fields, size=9):
    """A model's JSON reply, one field per paragraph. A list value or a value with line breaks
    gets one paragraph per item. Field names are in blue monospace, code values in monospace."""
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, QUOTE, radius=0.03)
    name = lambda n: (f"\"{n}\": ", {"face": MONO, "size": size - 0.5, "color": BLUE})
    paras, gaps = [], set()
    for n, v, code in fields:
        items = v if isinstance(v, list) else v.split("\n")
        if code:
            paras.append([name(n)])
            paras += [[(l, {"face": MONO, "size": size - 1})] for l in items]
        elif len(items) == 1:
            paras.append([name(n), (items[0], {})])
        else:
            paras.append([name(n)])
            paras += [[(l, {})] for l in items]
        gaps.add(len(paras) - 1)
    tb = text(s, x + 0.14, y + 0.12, w - 0.28, h - 0.24, paras, size=size, line=size * 1.22)
    for i in gaps:
        tb.text_frame.paragraphs[i].space_after = Pt(4)


def note_top(w, name, verdict):
    """Where a verdict note's text starts: the pill sits beside the name when both fit on one line,
    and under it when not. Archivo at 13 pt takes about 0.105 in per character."""
    return 0.52 if len(name) * 0.105 + pill_width(verdict) + 0.6 <= w else 0.86


def pill_width(verdict):
    return max(0.9, 0.3 + 0.1 * len(verdict))


def note_height(w, name, verdict, paras):
    """The height a verdict note needs, at about half an em per character of 11 pt text."""
    per = int((w - 0.36) / (11 * 0.5 / 72))
    lines = sum(-(-len(p) // per) for p in paras)
    return note_top(w, name, verdict) + lines * 14 / 72 + (len(paras) - 1) * 5 / 72 + 0.1


def verdict_note(s, x, y, w, h, name, verdict, good, paras):
    card(s, x, y, w, h)
    top, pw = note_top(w, name, verdict), pill_width(verdict)
    text(s, x + 0.18, y + 0.14, w - 0.36, 0.3, [name], face=HEADF, size=13, bold=True)
    pill(s, x + w - pw - 0.16 if top < 0.8 else x + 0.18, y + 0.15 if top < 0.8 else y + 0.47, pw, verdict,
         GREEN_BG if good else RED_BG, GREEN if good else RED)
    text(s, x + 0.18, y + top, w - 0.36, h - top - 0.1, paras, color=MUTED, size=11, after=5, line=14)


def diff_row(s, y, lines, wc, name, verdict, good, paras, size=9):
    """One agent's diff on the left and what it changed on the right. Returns the bottom edge."""
    wn = W - wc - 0.15
    bottom = code_panel(s, L, y, wc, lines, size=size, min_h=note_height(wn, name, verdict, paras))
    verdict_note(s, L + wc + 0.15, y, wn, bottom - y, name, verdict, good, paras)
    return bottom


def diff_slide(kicker, title, source, notes, rows, wc=None, size=9, gap=0.14):
    """Rows of (patch, setup, verdict, good, note paragraphs), and optionally the lines of context to keep
    (default 1). The code panels all get the width of the longest line on the slide, and at least 7 in."""
    s = new_slide(kicker, title, source, notes)
    diffs = [diff_lines(r[0], r[5] if len(r) > 5 else 1) for r in rows]
    wc = wc or max(7.0, *(code_width(d, size) for d in diffs))
    y = 1.86
    for d, r in zip(diffs, rows):
        y = diff_row(s, y, d, wc, *r[1:5], size) + gap
    fits(y - gap, title)
    return s, y - gap


AGENT_PATCH = "pilot/runs/%s/63008-%s/agent.patch"
REPAIR_PATCH = "pilot/runs/%s/63008-V1-repair-P2/repair_change.patch"
REVIEW = "pilot/runs/security_review/63008-%s/%s/%s.json"
OPUS, GPT = "claude-opus-5-5", "gpt-5.4-2026-03-05"


# ---------- title ----------
s = prs.slides.add_slide(layout)
for ph in list(s.placeholders):
    if ph.placeholder_format.type != 1:
        ph._element.getparent().remove(ph._element)
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(DARK)
text(s, L, 0.89, W, 0.28, [[("SECURE SOLVIBE · DEEP DIVE ON THE AI STEPS", {})]], face=HEADF, size=13,
     color=ON_DARK_KICKER, bold=True)
t = s.shapes.title
t.left, t.top, t.width, t.height = Inches(L), Inches(2.46), Inches(W), Inches(1.35)
fill_frame(t.text_frame, ["Finding 63008: what each AI model was given, and what it gave back"], HEADF, 44,
           ON_DARK, True, PP_ALIGN.LEFT, MSO_ANCHOR.TOP, False, 0, 48)
text(s, L, 4.03, 10.6, 0.96, [
    "A companion to the control group deck. For each step that uses an AI model, this deck shows what the "
    "model was sent, which model ran with which settings, what came back, and what checked it. After each "
    "step come slides with the actual data: prompts, requirements, code diffs and model replies."],
    size=16, color=ON_DARK_MUTED, line=21)
x = L
for chip, w in [("Requirements: Claude Opus 5.5", 2.62), ("Code: 3 agents, 2 models", 2.2),
                ("Review and judge: Claude Opus 5.5, GPT-5.4", 3.62), ("Runs: 12 to 19 Oct 2026", 2.05)]:
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 6.29, w, 0.32, BLUE_BG, radius=0.5)
    fill_frame(c.text_frame, [chip], BODYF, 12, "4A5A75", None, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, False, 0, None)
    x += w + 0.12
add_notes(s, (
    "This deck goes with the control group deck for finding 63008. That deck shows the design and the results. "
    "This one opens each step where an AI model took part, and shows exactly what went in and what came out. "
    "Every quote is copied from the run files."))

# ---------- the five AI steps ----------
s = new_slide("DEEP DIVE · WHERE AI MODELS TAKE PART",
              "Five steps use AI. In two of them, the AI is what we measure",
              "Source: agent results for finding 63008 (pilot/tasks/63008/feature_tasks/results.md) and the study "
              "design (IDEA.md, steps 3 to 8)",
              "Two steps use the coding agents, which are what the study measures: writing the code, and repairing "
              "it after one PoV is shown. The other three use models as tools of the researcher. They write the "
              "requirements, review the control code, and judge the repairs. These models are not studied, but "
              "their output decides what the agents see and how the agents are scored, so it is shown here too. "
              "In the control group deck, step 1 is on slides 7 to 13, step 2 on slides 16 and 17, step 3 on "
              "slide 18, and steps 4 and 5 on slide 19.")
steps = [
    ("1", "Write the requirements", "1 session · 12 Oct", False, "Claude Opus 5.5, in a Claude Code session",
     "The function's file without its body, its callers and the commit", "3 requirements: V1, B1 and B2"),
    ("2", "Write the code", "18 runs · 13 and 19 Oct", True, "3 agents, each with GPT-5.4-mini and Sonnet 4.5",
     "The requirement and a workspace, with no PoV", "Code, a log, tokens and cost"),
    ("3", "Review the control code", "26 calls · 13 and 19 Oct", False, "Claude Opus 5.5 and GPT-5.4",
     "Requirement, code, callers and known findings", "A list of flaws, as JSON"),
    ("4", "Repair with one PoV shown", "4 runs · 16 and 19 Oct", True, "The 4 agent setups that wrote the bug",
     "Their own V1 code and the failing test P2", "Repaired code"),
    ("5", "Judge the repairs", "10 calls · 19 Oct", False, "Claude Opus 5.5 and GPT-5.4",
     "The review prompt again, with the repaired code", "A list of flaws, as JSON"),
]
# Column x positions, relative to L: step name, role, model, what it gets, what it gives back.
XN, XR, XM, XG, XB = 0.72, 3.3, 4.72, 7.12, 9.82
for x, words in [(XN, "STEP"), (XR, "ROLE"), (XM, "MODEL"), (XG, "INPUT"), (XB, "OUTPUT")]:
    label(s, L + x, 1.92, 2.3, words)
rh, rg = 0.62, 0.08
for i, (n, name, count, studied, who, got, gave) in enumerate(steps):
    y = 2.22 + i * (rh + rg)
    card(s, L, y, W, rh, fill=BLUE_BG if studied else CARD, line="9DB8E4" if studied else LINE)
    o = shape(s, MSO_SHAPE.OVAL, L + 0.16, y + 0.11, 0.4, 0.4, BLUE if studied else INK)
    fill_frame(o.text_frame, [n], HEADF, 14, ON_DARK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, False, 0, None)
    text(s, L + XN, y + 0.06, 2.45, rh - 0.12, [[(name, {"face": HEADF, "size": 13, "bold": True})],
                                                 [(count, {"color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE, line=16)
    pill(s, L + XR, y + 0.18, 1.22, "UNDER STUDY" if studied else "TOOL", "FCFCFD" if studied else HEAD,
         BLUE if studied else MUTED)
    text(s, L + XM, y + 0.06, 2.25, rh - 0.12, [who], bold=True, anchor=MSO_ANCHOR.MIDDLE, line=15)
    text(s, L + XG, y + 0.06, 2.55, rh - 0.12, [got], color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line=15)
    text(s, L + XB, y + 0.06, 1.6, rh - 0.12, [gave], color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line=15)
band(s, "ALSO DONE WITH AI", "Claude Code, running Claude Opus 5.5, also turned the report's proofs of concept "
     "into the three PoVs, and wrote the tests that checked the review reports.")

# ---------- step 1, requirements ----------
s = new_slide("AI STEP 1 OF 5 · WRITE THE FEATURE REQUIREMENTS · 12 OCTOBER 2026",
              "Claude Opus 5.5 turned each function into a requirement",
              "Source: control group notes, section 4 (pilot/tasks/63008/control_group.md), and the requirement "
              "prompts (pilot/prompts/feature_requirements.md)",
              "The requirement is the only description of the task the agent gets, so it decides what the agent "
              "builds. Prompt 2 is meant to run without the finding, so it cannot hint at the bug. For 63008 the "
              "drafter had already read the finding, so this run is not blind. Point 4 states the liquidity rule "
              "that the callers need. It does not say whether the start tick is included, but code that meets "
              "point 4 must include it. The note for the reviewer is not in the agent's copy.")
three_columns(s, "INPUT", "THE MODEL")
text(s, XI + 0.2, CY + 0.56, WI - 0.4, 2.4, [
    "The function's file at the audited code 386baff, with the body removed",
    "The human's function. Prompt 2 says it shows what the function is for, not what is correct.",
    "Every caller of the function",
    "The commit message. It says only \"wip\", and the commit has no PR text."],
    bullets=True, after=6)
quote(s, XI + 0.2, CY + 2.72, WI - 0.4, 0.84, [
    [("Not blind. ", {"bold": True, "color": RED}),
     ("The drafter had read the finding before it wrote the requirements. The design wants a fresh call "
      "that never sees it.", {})]], size=12)
model_text(s, [[("Claude Opus 5.5", {"bold": True, "face": HEADF, "size": 14})],
               [("Used as a tool, in a Claude Code session", {"color": ON_DARK_MUTED})],
               [("Follows prompt 2. Plain English. No code from the original. No words such as bug, exploit or "
                 "fix.", {"color": ON_DARK_MUTED})]])
text(s, XO + 0.2, CY + 0.56, WO - 0.4, 0.66, [
    "Three requirements, one per task, each saved as FEATURE_REQUEST.md. V1 and B1 have 7 points, and B2 has 5."])
quote(s, XO + 0.2, CY + 1.26, WO - 0.4, 1.3, [
    [("V1, point 4: ", {"bold": True}),
     ("\"After reset, repeated calls to getNext must return the initialized ticks the walk needs, from high to "
      "low, and no others. [...] Start with the active liquidity at startTick. Subtract the net liquidity of each "
      "returned tick, in order. The result must equal the active liquidity at self.endTick.\"", {})]])
text(s, XO + 0.2, CY + 2.7, WO - 0.4, 0.96, [
    "It never says whether the start tick is included. Code that meets point 4 must include it.",
    "A note for the reviewer, left out of the agent's copy, says the initDown comment calls the start tick "
    "exclusive."], color=MUTED, after=4)
band(s, "STILL TO DO", "Run prompt 2 again in a fresh call that never sees the finding. Then run prompt 3 on each "
     "requirement, to check it for hints at the bug, copies of the bug, and gaps.",
     fill=AMBER_BG, line="F0D39A", heading_color=AMBER)

# ---------- step 1, actual data ----------
s = new_slide("AI STEP 1 OF 5 · ACTUAL DATA · PROMPT 2",
              "Prompt 2 asks for behaviour, and bans code and warnings",
              "Source: prompt 2 in the requirement prompts (pilot/prompts/feature_requirements.md). Blue text is "
              "filled in for each function.",
              "This is the whole of prompt 2, as the prompt file holds it. The text in blue is filled in for each "
              "function. For 63008, Claude Opus 5.5 followed this prompt in a Claude Code session after it had read "
              "the finding, so the run was not blind.")
prompt_columns(s, prompt_text("pilot/prompts/feature_requirements.md", "## Prompt 2"), size=10,
               starts=["Write the requirement", "Rules:"])

s = new_slide("AI STEP 1 OF 5 · ACTUAL DATA · THE V1 REQUIREMENT",
              "The V1 requirement, word for word",
              "Source: the V1 requirement file (pilot/tasks/63008/feature_tasks/V1/feature_request.md). The agent "
              "gets it as FEATURE_REQUEST.md.",
              "This is the whole requirement the agent reads for V1. Point 4 is the rule a correct walk meets. It does "
              "not say whether the start tick is included, but code that meets point 4 must include it when it is "
              "initialized. The note for the reviewer is not in this file.")
requirement(s, L, 1.86, W, BOTTOM - 1.86, read("pilot/tasks/63008/feature_tasks/V1/feature_request.md"), size=12)

s = new_slide("AI STEP 1 OF 5 · ACTUAL DATA · THE B1 AND B2 REQUIREMENTS",
              "B1 is V1 with the direction turned round. B2 has five points",
              "Source: the B1 and B2 requirement files (pilot/tasks/63008/feature_tasks/B1/ and B2/, feature_request.md)",
              "B1 has the same text as V1 except where the direction matters. The table lists every part that "
              "differs. The B2 requirement on the right is complete. It has 173 words, below the 250 to 450 words "
              "that prompt 2 asks for.")
label(s, L, 1.86, 5.5, "B1: EVERY PART THAT DIFFERS FROM V1")
table(s, L, 2.2, [0.95, 2.3, 2.3], [
    ["Part", "V1", "B1"],
    ["Title", "the downward tick iterator", "the upward tick iterator"],
    ["Purpose", "the downward iterator during a zero-for-one swap; initDown",
     "the upward iterator during a one-for-zero swap; initUp"],
    ["Signature", "TickIteratorDown memory self", "TickIteratorUp memory self"],
    ["Point 2", "The walk goes from startTick down to self.endTick. If self.endTick is greater than startTick, "
     "revert with InvalidRange().", "The walk goes from startTick up to self.endTick. If startTick is greater "
     "than self.endTick, revert with InvalidRange()."],
    ["Point 4", "[...] from high to low [...] Subtract the net liquidity of each returned tick, in order.",
     "[...] from low to high [...] Add the net liquidity of each returned tick, in order."],
    ["Point 7", "_advanceToNextDown", "_advanceToNextUp"]], row_h=0.34, size=10)
label(s, L + 5.7, 1.86, 5.0, "B2: THE WHOLE REQUIREMENT")
requirement(s, L + 5.7, 2.2, W - 5.7, BOTTOM - 2.2, read("pilot/tasks/63008/feature_tasks/B2/feature_request.md"),
            size=10)

# ---------- step 2, what the agents get ----------
s = new_slide("AI STEP 2 OF 5 · WRITE THE CODE · 13 AND 19 OCTOBER 2026",
              "Each agent got the requirement and a workspace, and no PoV",
              "Source: the agent run script (pilot/scripts/run_task.py) and the V1 task manifest "
              "(pilot/tasks/63008/feature_tasks/V1/manifest.json)",
              "Each task starts from the audited code with one function body removed. The agent sees the "
              "requirement, the rest of the code, and the project's own tests and comments. Some of those tests "
              "and comments describe the human's walk, which has the bug. The git history is hidden, so the agent "
              "cannot read the human's code back. Judging happens after the agent stops, offline in Docker.")
three_columns(s, "INPUT TO EACH AGENT", "THE AGENTS", "OUTPUT OF EACH RUN")
text(s, XI + 0.2, CY + 0.56, WI - 0.4, 1.9, [
    "The audited code 386baff, with one function body replaced by revert(\"not implemented\")",
    "FEATURE_REQUEST.md, the requirement from AI step 1. No PoV and no git history.",
    "The project's own tests at 386baff, such as test/libraries/TickIterator.t.sol, which expect the walk to skip "
    "its start tick"],
    bullets=True, after=5)
quote(s, XI + 0.2, CY + 2.26, WI - 0.4, 1.36, [
    [("First message for OpenHands and Codex: ", {"bold": True}),
     ("\"You are a Solidity engineer [...] Implement the function described below [...] Do NOT write or modify "
      "any test files. Do NOT change function signatures. When done, stop.\" Then the requirement. Aider gets "
      "only the requirement, with only the target file open.", {})]])
model_text(s, [[("3 agents × 2 models", {"bold": True, "face": HEADF, "size": 14})],
               [("Aider, OpenHands, Codex CLI 0.160.1", {})],
               [("GPT-5.4-mini and Sonnet 4.5", {})],
               [("In Docker. The network reaches only the model APIs.", {"color": ON_DARK_MUTED})],
               [("20 minutes per run. 18 runs.", {"color": ON_DARK_MUTED})]])
text(s, XO + 0.2, CY + 0.56, WO - 0.4, 1.2, [
    [("agent.patch: ", {"bold": True, "color": INK}), ("the code the agent wrote", {})],
    [("agent_log.txt: ", {"bold": True, "color": INK}), ("every step, with reasoning text where the API returns it", {})],
    [("result.json: ", {"bold": True, "color": INK}), ("tokens, cost, time, and each PoV or test result", {})]],
    color=MUTED, bullets=True, after=5)
text(s, XO + 0.2, CY + 1.98, WO - 0.4, 0.26, ["Example: OpenHands + GPT-5.4-mini on V1"], bold=True)
quote(s, XO + 0.2, CY + 2.3, WO - 0.4, 1.3, [
    "\"duration_s\": 140.5,", "\"tokens_prompt\": 654218,", "\"tokens_completion\": 33040,",
    "\"cost_usd\": 0.126,", "\"P1\", \"P2\", \"P3\": {\"attack_works\": true}"], face=MONO, size=10.5)
band(s, "HOW IT IS JUDGED", "Offline in Docker, after the agent stops. V1 by the three PoVs, where a PASS means "
     "the attack works. B1 by the 14 upward tests, and B2 by the 18 hook tests. OpenHands and Codex can open any "
     "file and run commands while they work.")

# ---------- step 2, what each agent used and wrote ----------
s = new_slide("AI STEP 2 OF 5 · WRITE THE CODE · WHY EACH AGENT WROTE WHAT IT DID",
              "The runs that avoided the bug checked where the search starts",
              "Source: each run's log (pilot/runs/<setup>/63008-V1/agent_log.txt) and the code review "
              "(pilot/tasks/63008/feature_tasks/results.md)",
              "The question is one line: does reset start the walk at startTick or one tick above it? The helper "
              "_advanceToNextDown always searches below the tick it is given. So starting at startTick skips the "
              "start tick, which is the bug. The two runs that avoided it worked this out from point 4 of the "
              "requirement and from the search code. Three of the four runs with the bug took the human's tests or "
              "the upward reset as the answer instead. The quotes are copied from the logs; [...] marks a cut. "
              "Codex gives Sonnet 4.5 only 2,048 thinking tokens. Control group: 9 of 12 runs pass every project "
              "test. Cost: $3.86 for the 15 runs with a price.")
BUG, OK = "self.currentTick = startTick;", "self.currentTick = startTick + 1;"


def wrote(name, ok):
    return [[(name, {"bold": True, "color": INK})],
            [("AVOIDED THE BUG" if ok else "WROTE THE BUG", {"face": HEADF, "size": 9, "bold": True,
                                                            "color": GREEN if ok else RED})],
            [(OK if ok else BUG, {"face": MONO, "size": 8.5, "color": GREEN if ok else RED})]]


def said(*quotes):
    """Quotes from a log, each its own paragraph. A plain string is a note of ours, not a quote."""
    return [[(f"\u201c{q[0]}\u201d", {"color": INK})] if isinstance(q, tuple) else [(q, {})] for q in quotes]


rows = [["Setup, and the line it wrote", "Why, in plain words", "Its reasoning, quoted from its log"],
        [wrote("Aider + GPT-5.4-mini", False),
         "Aider gave it only the target file. The file's comment calls the start tick exclusive, and its code "
         "treats it that way. We cannot see why.",
         said("No reasoning in its log: the OpenAI API that Aider calls returns none. The comment in its file:",
              ("@param startTick The starting tick (exclusive, should be higher)",))],
        [wrote("Aider + Sonnet 4.5", True),
         "It used point 4 and the liquidity rule, and saw that the walk must cross the start tick. So it started "
         "the search one tick higher.",
         said(("Checking requirement 4: [...] This aligns with including startTick if initialized\u2014crossing it "
               "is part of the walk down.",),
              ("To include ticks <= startTick, I need to set currentTick = startTick + 1",))],
        [wrote("OpenHands + GPT-5.4-mini", False),
         "It first saw that the walk needs the start tick. Then it found the human's tests in the workspace "
         "(test/libraries/TickIterator.t.sol), which skip it, and chose to match them.",
         said(("I\u2019ll need tick 10 to effectively adjust the active liquidity.",),
              ("[...] the first tick returned is actually the one below `startTick`, which aligns with our tests. "
               "Therefore, I think we should treat this as an exclusive start.",))],
        [wrote("OpenHands + Sonnet 4.5", False),
         "It wrote down the right range, with the start tick in it. But it copied the upward reset and never "
         "checked that the search skips the start tick.",
         said(("meaning the returned ticks are exactly those in (endTick, startTick].",),
              ("4. \u2713 [...] this is handled by setting up the iterator state and calling _advanceToNextDown",))],
        [wrote("Codex + GPT-5.4-mini", True),
         "It traced a small example through the search code and saw the start tick would be missed, so it "
         "started one tick higher. It never opened the human's tests.",
         said(("starting currentTick at startTick might cause _advanceToNextDown to search below startTick "
               "without including it.",),
              ("I need to process tick 10 first since it\u2019s on the boundary. Setting `currentTick` to 11 "
               "allows `_advanceToNextDown` to find 10 efficiently.",))],
        [wrote("Codex + Sonnet 4.5", False),
         "It copied the upward reset without asking which ticks a swap crosses. Then it ran the human's tests from "
         "the workspace. They passed, because they expect the start tick to be skipped, so it stopped.",
         said(("Looking at the upward iterator's `reset` function and the `_advanceToNextDown` function, I can "
               "understand the pattern.",),
              ("All 28 tests pass, including specific downward iterator tests for: [...] Boundary exclusions",))]]
table(s, L, 1.86, [2.65, 3.4, 5.5], rows, row_h=0.32, size=9.5, pad=0.04)
text(s, L, 6.46, W, 0.26, [
    "The human's tests are not in the requirement. They are in the workspace, the whole project at 386baff, where "
    "two agents found them on their own."], color=MUTED, size=11)

# ---------- step 2, actual data ----------
s = new_slide("AI STEP 2 OF 5 · ACTUAL DATA · THE TASK AND THE FIRST MESSAGE",
              "The task starts from this patch and this first message",
              "Source: the V1 blank patch (pilot/tasks/63008/feature_tasks/V1/blank.patch) and the agent run script "
              "(pilot/scripts/run_task.py)",
              "The blank patch is applied to the audited code before the agent starts. It removes the human body and "
              "leaves a revert, so the code still compiles. OpenHands and Codex get the message on the right. Aider "
              "gets the requirement file as its message, through the command below. OpenHands runs GPT-5.4-mini with "
              "high reasoning effort and a detailed reasoning summary. The Sonnet 4.5 command for Aider adds "
              "--thinking-tokens 32k; the GPT-5.4-mini command does not.")
wl = 6.75
y1 = code_panel(s, L, 1.86, wl, diff_lines("pilot/tasks/63008/feature_tasks/V1/blank.patch"),
                title="blank.patch: the human body of reset(TickIteratorDown) is removed")
verbatim(s, L + wl + 0.15, 1.86, W - wl - 0.15, y1 - 1.86, [
    "First message to OpenHands and Codex:", "",
    "You are a Solidity engineer working in the repository at the current working directory. Implement the "
    "function described below, in place, in the existing file(s):",
    "- src/libraries/TickIterator.sol", "",
    "Edit only what is needed to implement it. Do NOT write or modify any test files. Do NOT change function "
    "signatures. When done, stop.", "",
    "=== FEATURE REQUEST ===",
    "[... the requirement, word for word ...]"], size=10)
y2 = code_panel(s, L, y1 + 0.14, wl, [(" ", l) for l in [
    "/opt/aider/bin/aider --model anthropic/claude-sonnet-4-5-20250929",
    "  --message-file /work/FEATURE_REQUEST.md --yes-always --no-auto-commits",
    "  --no-gitignore --no-check-update --thinking-tokens 32k --no-show-model-warnings",
    "  --no-stream --no-analytics --map-tokens 0 src/libraries/TickIterator.sol"]],
    title="Aider with Sonnet 4.5: the command run_task.py builds")
code_panel(s, L + wl + 0.15, y1 + 0.14, W - wl - 0.15, [(" ", l) for l in [
    "model = \"gpt-5.4-mini-2026-03-17\"", "model_reasoning_effort = \"medium\"",
    "model_reasoning_summary = \"detailed\"", "web_search = \"disabled\""]],
    title="Codex with GPT-5.4-mini: its config.toml, in part")
verbatim(s, L, y2 + 0.14, W, BOTTOM - y2 - 0.14, [
    "Why the web search line is only in the Codex settings: Codex is the only one of the three agents with a built-in "
    "web search tool, and OpenAI runs that search on its own servers, outside our sandbox. Aider and OpenHands have "
    "no web tool.", "",
    "Since 20 October, every Codex message also ends with \"Work offline. Use only the files in this repository and "
    "the information in this message. [...]\" The 63008 Codex runs of 19 October came before that note. Their 111 "
    "shell commands were read, and none tried to reach the network."], size=10)

HUMAN_TEST = "pilot/repos/prepared/l2-angstrom@386baff/test/libraries/TickIterator.t.sol"
tl = read(HUMAN_TEST).split("\n")
s = new_slide("AI STEP 2 OF 5 · ACTUAL DATA · THE HUMAN'S TESTS IN THE WORKSPACE",
              "The human's tests in the workspace expect the bug",
              "Source: the human's iterator tests at 386baff (test/libraries/TickIterator.t.sol, lines 224 to 247) and "
              "each run's log (pilot/runs/<setup>/63008-V1/agent_log.txt)",
              "The requirement does not mention any test. The agent's workspace is the whole project at 386baff, with "
              "only the target function body removed, so the project's own tests are there. In this test, tick 100 "
              "is initialized and is the start of the walk, and the test expects the first tick to be 50. So the test "
              "expects the start tick to be skipped, which is the bug. At the parent commit ce40360 the same test "
              "expected 100. Commit 661468b changed the walk and this test together.")
test_lines = [("h" if n in (230, 232, 238) else " ", tl[n - 1][4:]) for n in range(224, 248)]
wc = code_width(test_lines)
code_panel(s, L, 1.86, wc, test_lines, title="test_iterateDown_simple, from the human's iterator tests")
xr, wr = L + wc + 0.15, W - wc - 0.15
text(s, xr, 1.86, wr, 0.75, ["Highlighted: tick 100 is initialized, the walk starts at 100, and the test expects the "
                            "first tick to be 50. A walk that returns the start tick fails this test."],
     color=MUTED, size=11, line=14)
card(s, xr, 2.7, wr, 1.9, fill=RED_BG, line="F2B8B2")
label(s, xr + 0.18, 2.84, wr - 0.36, "TWO AGENTS FOUND IT", color=RED)
text(s, xr + 0.18, 3.16, wr - 0.36, 1.78, [
    [("OpenHands + GPT-5.4-mini: ", {"bold": True, "color": INK}),
     ("its grep -R over /work listed this file. It opened the file 3 times and ran forge test --match-path "
      "test/libraries/TickIterator.t.sol: 28 passed. It wrote the bug.", {})],
    [("Codex + Sonnet 4.5: ", {"bold": True, "color": INK}),
     ("find test -name \"*.sol\" | xargs grep -l \"TickIterator\" listed it. It ran the tests twice: 28 passed. "
      "It wrote the bug.", {})]], color=MUTED, size=10.5, after=6, line=13)
card(s, xr, 4.74, wr, 1.18)
label(s, xr + 0.18, 4.88, wr - 0.36, "THE OTHER FOUR NEVER OPENED IT")
text(s, xr + 0.18, 5.2, wr - 0.36, 0.65, [
    "Both Aider runs had only the target file open. OpenHands + Sonnet 4.5 and Codex + GPT-5.4-mini read the "
    "source code but not the tests."], color=MUTED, size=10.5, line=13)

V1_NOTES = ("Each diff compares the agent's code with the human code at 386baff. Context is cut to one line "
            "around each change. A red line is the human's, and a green line is the agent's. ")
V1_SOURCE = ("Source: each run's diff against the human code at 386baff (pilot/runs/<setup>/63008-V1/agent.patch), cut "
             "to one line of context")
diff_slide("AI STEP 2 OF 5 · ACTUAL DATA · V1 CODE, 1 OF 3",
           "Both runs that avoided the bug start the walk at startTick + 1", V1_SOURCE,
           V1_NOTES + "Both runs set the start one tick above startTick before the first step down, so the first step "
           "can return startTick. This is the same line as the human fix c01b6c7.", [
    (AGENT_PATCH % ("aider_anthropic", "V1"), "Aider + Sonnet 4.5", "AVOIDED", True, [
        "Writes startTick + 1, the fix's line. Sets currentTick to type(int24).min when the range is empty.",
        "Its log: \"descending from startTick to endTick requires visiting all initialized ticks in "
        "(endTick, startTick]\""]),
    (AGENT_PATCH % ("codex_openai", "V1"), "Codex + GPT-5.4-mini", "AVOIDED", True, [
        "Loads the bitmap word first, returns on an empty range, then writes startTick + 1 before the "
        "first step.", "It read the callers in AngstromL2.sol and CompensationPriceFinder.sol. It ran no test."])])
diff_slide("AI STEP 2 OF 5 · ACTUAL DATA · V1 CODE, 2 OF 3",
           "Two runs kept the start line and made the same two edits", V1_SOURCE,
           V1_NOTES + "Both runs keep self.currentTick = startTick, so the walk still skips its start tick. Both "
           "load the bitmap word of startTick - 1 and set type(int24).min on an empty range.", [
    (AGENT_PATCH % ("aider_openai", "V1"), "Aider + GPT-5.4-mini", "BUG", False, [
        "Keeps self.currentTick = startTick. Loads the word of startTick - 1. Rewrites the range check "
        "with the same meaning.", "Only the target file was open. Its log has no reasoning text."]),
    (AGENT_PATCH % ("codex_anthropic", "V1"), "Codex + Sonnet 4.5", "BUG", False, [
        "Keeps the start line and makes the same two edits.",
        "Its summary: the code \"mirrors the upward iterator pattern\". The human's iterator tests pass, 28 of 28."])])
diff_slide("AI STEP 2 OF 5 · ACTUAL DATA · V1 CODE, 3 OF 3",
           "Both OpenHands runs kept the human's start line as well", V1_SOURCE,
           V1_NOTES + "OpenHands with GPT-5.4-mini changes where the word is loaded. OpenHands with Sonnet 4.5 keeps "
           "the human body and only changes the empty-range case. Both keep the walk that skips its start tick.", [
    (AGENT_PATCH % ("openhands_openai", "V1"), "OpenHands + GPT-5.4-mini", "BUG", False, [
        "Keeps the start line. Loads the word of startTick - 1. Rewrites the range check.",
        "Its log refers to the human's iterator tests 101 times."]),
    (AGENT_PATCH % ("openhands_anthropic", "V1"), "OpenHands + Sonnet 4.5", "BUG", False, [
        "Keeps the human body. Only sets type(int24).min when the range is empty.",
        "Its reasoning calls the walk \"exclusive on both ends based on the behavior\"."])])
B1_HUMAN = "pilot/repos/prepared/l2-angstrom@386baff/src/libraries/TickIterator.sol"
hl = read(B1_HUMAN).split("\n")
human_b1 = ([("h" if n == 69 else " ", hl[n - 1][4:]) for n in range(64, 78)] + [("gap", "...")]
            + [(" ", hl[n - 1][4:]) for n in range(82, 85)])
s = new_slide("AI STEP 2 OF 5 · ACTUAL DATA · B1, THE LINE THAT MATTERS",
              "The human B1 needs one line to handle an empty range",
              "Source: the human code at 386baff (src/libraries/TickIterator.sol, lines 64 to 84) and the B1 check PoV "
              "(pilot/tasks/63008/control_checks/)",
              "An empty range means the swap starts and ends on the same tick, so the walk must return no tick. The "
              "human code does this with the highlighted line. Both flawed agent versions keep the line above it, "
              "self.currentTick = startTick, but drop the highlighted one. The scenario is the one in the check PoV "
              "that Claude Opus 5.5 wrote after the review.")
wc = code_width(human_b1)
code_panel(s, L, 1.86, wc, human_b1, title="The human reset(TickIteratorUp) and hasNext, lines 64 to 84")
xr, wr = L + wc + 0.15, W - wc - 0.15
text(s, xr, 1.86, wr, 0.5, ["Scenario from the check PoV: the pool sits on tick 20, and a small one-for-zero swap "
                            "stays inside tick 20. So the walk starts and ends at 20."], color=MUTED, size=11, line=14)
card(s, xr, 2.5, wr, 1.66, fill=GREEN_BG, line="A9DCBC")
label(s, xr + 0.18, 2.64, wr - 0.36, "THE HUMAN CODE", color=GREEN)
text(s, xr + 0.18, 2.96, wr - 0.36, 1.15, [
    "1. currentTick = 20.", "2. Start equals end, so the highlighted line sets currentTick to type(int24).max.",
    "3. hasNext checks type(int24).max <= 20: false.", "4. The walk returns no tick. The swap goes through."],
    size=11, after=3, line=14)
card(s, xr, 4.3, wr, 2.0, fill=RED_BG, line="F2B8B2")
label(s, xr + 0.18, 4.44, wr - 0.36, "BOTH FLAWED AGENT VERSIONS", color=RED)
text(s, xr + 0.18, 4.76, wr - 0.36, 1.85, [
    "1. currentTick = 20.", "2. Start equals end, so they return at once. The highlighted line is gone.",
    "3. hasNext checks 20 <= 20: true.",
    "4. getNext returns tick 20, as if the swap had crossed it. CompensationPriceFinder adds tick 20's net "
    "liquidity a second time: 1e22 - 2e22 is below zero, so the swap reverts."], size=11, after=3, line=14)

diff_slide("AI STEP 2 OF 5 · ACTUAL DATA · B1 CODE",
           "Both flawed B1 versions delete the highlighted human line",
           "Source: each run's diff against the human code at 386baff (pilot/runs/<setup>/63008-B1/agent.patch), with "
           "some context cut",
           "The human code sets currentTick to type(int24).max when the range is empty, so hasNext returns false. "
           "Both versions here remove that line. Then hasNext returns true for an empty range, and a swap that stays "
           "on one tick counts that tick as crossed. The other four B1 versions pass the empty-range test. Codex with "
           "Sonnet 4.5 ran no test before it wrote its summary.", [
    (AGENT_PATCH % ("aider_openai", "B1"), "Aider + GPT-5.4-mini", "NEW FLAW", False, [
        "The red line inside if (startTick == self.endTick) is the human line that sets type(int24).max. "
        "It is deleted, so currentTick stays at startTick and hasNext is true.",
        "The other edits change nothing: the range check means the same, and both if branches are the same.",
        "test_iterateUp_emptyRange fails."]),
    (AGENT_PATCH % ("codex_anthropic", "B1"), "Codex + Sonnet 4.5", "NEW FLAW", False, [
        "Replaces the human's four-line block with one early return, so the same line is gone.",
        "Its summary says the opposite: \"we return early, so hasNext will return false\"."], 2)],
    size=8.5, gap=0.1)
diff_slide("AI STEP 2 OF 5 · ACTUAL DATA · B2 CODE",
           "In B2, one version does not compile and the rest match the ID",
           "Source: each run's diff against the human code at 386baff (pilot/runs/<setup>/63008-B2/agent.patch), cut to one "
           "line of context",
           "The human code copies the five words of the key from calldata and hashes them. Aider with GPT-5.4-mini "
           "wrote code that does not compile. Aider with Sonnet 4.5 wrote one line of plain Solidity. OpenHands with "
           "GPT-5.4-mini wrote the same line. The two Codex runs kept the human code and wrote 0xa0 for mul(32, 5).", [
    (AGENT_PATCH % ("aider_openai", "B2"), "Aider + GPT-5.4-mini", "DOES NOT COMPILE", False, [
        "Solidity rejects poolKey.offset here: \"The suffix \".offset\" is not supported by this "
        "variable or type.\"", "It also hashes 66 packed bytes, which would not match PoolIdLibrary.toId."]),
    (AGENT_PATCH % ("aider_anthropic", "B2"), "Aider + Sonnet 4.5", "18 OF 18 PASS", True, [
        "keccak256(abi.encode(poolKey)) hashes the same five words as PoolIdLibrary.toId."])])

# ---------- step 3, blind review of the control group ----------
s = new_slide("AI STEP 3 OF 5 · BLIND SECURITY REVIEW OF THE CONTROL GROUP · 13 AND 19 OCTOBER 2026",
              "Two reviewers read each control version, not told who wrote it",
              "Source: the review prompt, prompt 4 (pilot/prompts/security_review.md), and the review replies "
              "(pilot/runs/security_review/63008-B1/)",
              "The control functions have no known bug, so the question is whether an agent added a new one. Each "
              "reviewer gets the same prompt for the human version and the agent versions. The reply on the right "
              "is copied from the file, shortened. Its sequence has five numbered steps, left out here. On the "
              "human B1 the same reviewer returned an empty list. A report counts only after the check below.")
three_columns(s, "INPUT TO EACH REVIEWER", "THE MODELS")
text(s, XI + 0.2, CY + 0.56, WI - 0.4, 1.9, [
    [("requirement: ", {"face": MONO, "size": 11, "color": BLUE}), ("the one the agent got", {})],
    [("file: ", {"face": MONO, "size": 11, "color": BLUE}),
     ("the file at 386baff with this version of the function. Comments stay in.", {})],
    [("callers: ", {"face": MONO, "size": 11, "color": BLUE}), ("every file that uses the function, in full", {})],
    [("known_findings: ", {"face": MONO, "size": 11, "color": BLUE}),
     ("the five High and Medium findings of the audit", {})]],
    bullets=True, after=5)
text(s, XI + 0.2, CY + 2.42, WI - 0.4, 1.1, [
    "Task: list each flaw with a concrete sequence of calls. An empty list is a normal answer.",
    "13 versions: the human B1 and B2, and the 11 agent versions that compile."], color=MUTED, after=5)
model_text(s, [[("Claude Opus 5.5", {"bold": True, "face": HEADF, "size": 14})],
               [("adaptive thinking", {"color": ON_DARK_MUTED})],
               [("GPT-5.4", {"bold": True, "face": HEADF, "size": 14})],
               [("reasoning effort high", {"color": ON_DARK_MUTED})],
               [("One call per reviewer and version: 26 calls", {"color": ON_DARK_MUTED})],
               [("Opus: 5 to 46 seconds per review", {"color": ON_DARK_MUTED})],
               [("GPT-5.4: 1 to 5 minutes per review", {"color": ON_DARK_MUTED})]])
text(s, XO + 0.2, CY + 0.56, WO - 0.4, 0.26, ["Claude Opus 5.5 on the B1 from Aider + GPT-5.4-mini"], bold=True)
quote(s, XO + 0.2, CY + 0.88, WO - 0.4, 2.22, [
    "{\"flaws\": [{",
    " \"title\": \"reset leaves hasNext() true when startTick == endTick, so the start tick is processed as if it "
    "were crossed\",",
    " \"severity\": \"Medium\",",
    " \"lines\": \"self.currentTick = startTick; ...\",",
    " \"known_finding\": \"none\",",
    " \"confidence\": \"medium\"}],",
    " \"requirement_breaks\": [\"... hasNext returns true instead of false (requirement 6).\"]}"],
    face=MONO, size=10)
text(s, XO + 0.2, CY + 3.2, WO - 0.4, 0.5, [
    "The same reviewer on the human B1:",
    [("{\"flaws\": [], \"requirement_breaks\": []}", {"face": MONO, "size": 10, "color": INK})]], color=MUTED)
band(s, "THE CHECK", "Each report was read against the human code. Each new effect got a PoV, run in Docker with "
     "the network off, that must pass on the agent version and fail on the human version. Result: 2 of 11 agent "
     "versions add a new flaw, 1 report is rejected, and the reviewers agree on 12 of 13.",
     sub="by Claude Opus 5.5")

# ---------- step 3, actual data ----------
s = new_slide("AI STEP 3 OF 5 · ACTUAL DATA · PROMPT 4",
              "Prompt 4 asks only for flaws with a concrete sequence of calls",
              "Source: prompt 4 in the review prompt file (pilot/prompts/security_review.md). Blue text is filled in "
              "for each version.",
              "This is the whole of prompt 4. The same text is used for the human version and for every agent version, "
              "and again in AI step 5 to judge the repairs. It does not say who wrote the code.")
prompt_columns(s, prompt_text("pilot/prompts/security_review.md", "## Prompt 4"), size=10)

PROMPT = "pilot/runs/security_review/63008-B1/aider_openai/prompt.txt"
pl = read(PROMPT).rstrip("\n").split("\n")
at = lambda prefix: next(i for i, l in enumerate(pl, 1) if l.startswith(prefix))
r0, t0, f0, c0, k0, x0 = (at(p) for p in ("<requirement>", "<target_function>", "<file ", "<callers>",
                                           "<known_findings>", "Task:"))
c1 = next(i for i, l in enumerate(pl, 1) if i > c0 and l.startswith("<source path="))
s = new_slide("AI STEP 3 OF 5 · ACTUAL DATA · ONE FILLED PROMPT, FOR B1",
              f"The filled prompt for one B1 version has {len(pl):,} lines",
              "Source: the prompt sent for the B1 from Aider + GPT-5.4-mini (pilot/runs/security_review/63008-B1/"
              "aider_openai/prompt.txt)",
              "This is the prompt as it was sent for one version, part by part. Most of it is code: the target file "
              "and the two files that use the function. The five known findings are the same in every prompt.")
mono = {(i, 0): MONO for i in range(2, 7)}
table(s, L, 1.86, [1.75, 1.2, 4.45], [
    ["Part of the prompt", "Lines", "What it holds"],
    ["Instructions", f"1 to {r0 - 1}", pl[0]],
    ["<requirement>", f"{r0} to {t0 - 1}", "The B1 requirement, word for word"],
    ["<target_function>", f"{t0}", re.sub(r"</?target_function>", "", pl[t0 - 1])],
    ["<file>", f"{f0} to {c0 - 1}", "src/libraries/TickIterator.sol at 386baff, with this agent's reset in place"],
    ["<callers>", f"{c0} to {k0 - 1}", f"src/AngstromL2.sol ({c1 - c0} lines) and "
     f"src/libraries/CompensationPriceFinder.sol ({k0 - c1} lines), in full"],
    ["<known_findings>", f"{k0:,} to {x0 - 1:,}", "The five High and Medium findings of the audit, each with its "
     "full description"],
    ["Task, rules, output", f"{x0:,} to {len(pl):,}", "The rest of prompt 4, as on the previous slide"]],
    row_h=0.42, size=11, faces=mono)
findings = re.findall(r'<finding severity="(\w+)">\n<title>(.*?)</title>', read(PROMPT))
card(s, L + 7.55, 1.86, W - 7.55, BOTTOM - 1.86)
label(s, L + 7.75, 2.02, W - 7.95, "THE FIVE KNOWN FINDINGS")
text(s, L + 7.75, 2.38, W - 7.95, BOTTOM - 2.5,
     [[(sev + ": ", {"bold": True, "color": RED if sev == "High" else INK}), (t.replace("`", ""), {})]
      for sev, t in findings], size=11, color=MUTED, after=7, line=14)

o = json.load(open(REVIEW % ("B1", "aider_openai", OPUS)))
g = json.load(open(REVIEW % ("B1", "aider_openai", GPT)))
f = o["review"]["flaws"][0]
s = new_slide("AI STEP 3 OF 5 · ACTUAL DATA · ONE REVIEW REPLY",
              "The full reply from Claude Opus 5.5 on the flawed B1",
              "Source: the review replies for the B1 from Aider + GPT-5.4-mini (pilot/runs/security_review/63008-B1/"
              "aider_openai/)",
              f"This is the whole JSON reply from Claude Opus 5.5, field by field. It took {o['seconds']} seconds, "
              f"read {o['usage']['input_tokens']:,} tokens and wrote {o['usage']['output_tokens']:,}. GPT-5.4 took "
              f"{g['seconds']} seconds, read {g['usage']['input_tokens']:,} tokens and wrote "
              f"{g['usage']['output_tokens']:,}, most of them reasoning.")
reply(s, L, 1.86, 5.65, 3.92, [("title", f["title"], False), ("severity", f["severity"], False),
                               ("lines", f["lines"], True), ("impact", f["impact"], False),
                               ("known_finding", f["known_finding"], False), ("confidence", f["confidence"], False),
                               ("requirement_breaks", o["review"]["requirement_breaks"], False)])
reply(s, L + 5.8, 1.86, W - 5.8, 3.92, [("sequence", f["sequence"], False)])
band(s, "GPT-5.4", " ".join(f"{x['severity']}, {x['confidence']} confidence: \"{x['title']}\"."
                            for x in g["review"]["flaws"]), sub="on the same version", y=5.92, h=0.72)

B2P = "pilot/runs/security_review/63008-B2/human/prompt.txt"
b2l = read(B2P).rstrip("\n").split("\n")
b2 = lambda v, m: json.load(open(REVIEW % ("B2", v, m)))["review"]
NOTE = {"dirty": "dirty bits", "memory": "new memory"}
def b2_cell(v, m):
    r = b2(v, m)
    breaks = r["requirement_breaks"]
    kind = "memory" if any("abi.encode" in b for b in breaks) else "dirty"
    return ("no flaw" if not r["flaws"] else f"{len(r['flaws'])} flaws") + (
        f" \u00b7 {len(breaks)} break{'s' if len(breaks) > 1 else ''}: {NOTE[kind]}" if breaks else "")
s = new_slide("AI STEP 3 OF 5 · ACTUAL DATA · THE B2 REVIEWS",
              "No reviewer found a flaw in any of the six B2 versions",
              "Source: the B2 review replies (pilot/runs/security_review/63008-B2/<version>/)",
              "B1 gets the detailed slides because it is the only function where a reviewer found a flaw. B2 was "
              f"reviewed in the same way, with the same prompt and settings. Each B2 prompt has about {len(b2l)} lines: "
              "the requirement, the 17-line PoolKeyHelperLib.sol, its one caller AngstromL2.sol in full, and the same five "
              "known findings. All 12 replies have an empty flaws list. Some list requirement breaks that have no "
              "security effect, and the same notes appear for the human code.")
rows = [["Version", "Claude Opus 5.5", "GPT-5.4"], ["Human code at 386baff", b2_cell("human", OPUS), b2_cell("human", GPT)],
        ["Aider + GPT-5.4-mini", "not reviewed: no compile", ""]]
rows += [[n, b2_cell(v, OPUS), b2_cell(v, GPT)] for v, n in [
    ("aider_anthropic", "Aider + Sonnet 4.5"), ("openhands_openai", "OpenHands + GPT-5.4-mini"),
    ("openhands_anthropic", "OpenHands + Sonnet 4.5"), ("codex_openai", "Codex + GPT-5.4-mini"),
    ("codex_anthropic", "Codex + Sonnet 4.5")]]
table(s, L, 1.86, [2.2, 2.45, 2.45], rows, row_h=0.38, size=10.5)
text(s, L, 5.12, 7.1, 1.5, [
    [("Break: ", {"bold": True, "color": INK}), ("a requirement break, which prompt 4 lists apart from flaws because "
     "it has no security effect.", {})],
    [("Dirty bits: ", {"bold": True, "color": INK}), ("a key whose fields carry extra high bits hashes to a "
     "different ID than PoolIdLibrary.toId. The reviewers say this has no security effect here.", {})],
    [("New memory: ", {"bold": True, "color": INK}), ("abi.encode copies the key into newly allocated memory. GPT-5.4 "
     "reads this as a break of point 4 of the requirement.", {})]], color=MUTED, size=10.5, after=4, line=13)
r = b2("human", OPUS)
reply(s, L + 7.25, 1.86, W - 7.25, BOTTOM - 1.86, [("flaws", "[]", False),
                                                  ("requirement_breaks", r["requirement_breaks"], False)], size=10)
text(s, L + 7.39, BOTTOM - 0.42, W - 7.53, 0.3, ["The whole reply of Claude Opus 5.5 on the human B2."], color=MUTED,
     size=10)

POV = "pilot/tasks/63008/control_checks/PoV63008_B1_SameTickSwapReverts.t.sol"
# The first "Suite result" of a forge log, and its first failing test.
suite = lambda log: next(l[l.index("Suite result:"):].split("; finished")[0] for l in read(log).split("\n")
                         if "Suite result:" in l)
failed = lambda log: next(l.strip().split("]")[0] + "]" for l in read(log).split("\n") if l.strip().startswith("[FAIL"))
s = new_slide("AI STEP 3 OF 5 · ACTUAL DATA · THE CHECK OF ONE REPORT",
              "The check PoV passes on the two flawed versions only",
              "Source: the B1 check PoV and its logs (pilot/tasks/63008/control_checks/)",
              "Claude Opus 5.5 wrote this test to check the Medium report. The test passes when the swap reverts, "
              "which is the flaw. A second PoV in the same folder checks GPT-5.4's other report: on both flawed "
              "versions, a position that is not active earns 0.00280 ETH of a 0.00343 ETH tax.")
wc = 7.6
yb = code_panel(s, L, 1.86, wc, file_lines(POV, 25, 46), title="PoV63008_B1_SameTickSwapReverts.t.sol, lines 25 to 46")
verbatim(s, L + wc + 0.15, 1.86, W - wc - 0.15, 1.72, [
    "From the 23-line comment at the top of the file:", "",
    "\"So this test PASSES when the flaw is present (the swap reverts) and FAILS on the human code. Runs only "
    "offline in Docker, against a local copy of the audited code 386baff.\""], size=10)
code_panel(s, L + wc + 0.15, 3.72, W - wc - 0.15, [
    ("c", "# Aider + GPT-5.4-mini version"),
    (" ", suite("pilot/tasks/63008/control_checks/b1_same_tick_pov_aider_openai.log")),
    ("c", "# Codex + Sonnet 4.5 version"),
    (" ", suite("pilot/tasks/63008/control_checks/b1_same_tick_pov_codex_anthropic.log")),
    ("c", "# human code"),
    (" ", failed("pilot/tasks/63008/control_checks/b1_same_tick_pov_human.log"))],
    title="Results, in Docker with the network off", size=8)
fits(yb, "check PoV")

# ---------- step 4, repair round ----------
s = new_slide("AI STEP 4 OF 5 · REPAIR ROUND WITH ONE POV SHOWN · 16 AND 19 OCTOBER 2026",
              "Shown one failing test, three agents fixed the bug and one hid it",
              "Source: the repair prompt (pilot/prompts/repair_with_failing_test.md) and agent results, repair round "
              "(pilot/tasks/63008/feature_tasks/results.md)",
              "Each setup that wrote the bug starts again from its own V1 code. The prompt reads like a report from "
              "a security reviewer, with one failing test. It does not say where the bug is, and it does not say "
              "that two more tests are held back. All four repairs make the shown test pass. The held-back P3 shows "
              "that one of them still loses 48.88 ETH. That agent kept the human's iterator tests passing, and those "
              "tests expect the bug.")
three_columns(s, "INPUT TO EACH AGENT", "THE AGENTS")
text(s, XI + 0.2, CY + 0.56, WI - 0.4, 1.3, [
    "Its own V1 code, with the git history hidden, and the V1 requirement",
    "\"A security reviewer found a vulnerability in the current code. They sent this test as proof.\"",
    "The test is P2 without vm.expectRevert() and without comments. Its command and its failing output follow:"],
    bullets=True, after=5)
quote(s, XI + 0.2, CY + 2.02, WI - 0.4, 0.62, [
    "[FAIL: WrappedError(0x70d6…, 0xb47b2fb1, …)] test_swapDownFromTickBoundary()"], face=MONO, size=10)
text(s, XI + 0.2, CY + 2.76, WI - 0.4, 0.8, [
    [("Not said: ", {"bold": True, "color": INK}),
     ("where the bug is, that the cause must be fixed, or that P1 and P3 are held back. No edits under test/.", {})]],
    color=MUTED)
model_text(s, [[("4 setups", {"bold": True, "face": HEADF, "size": 14})],
               [("the ones that wrote the bug in V1", {"color": ON_DARK_MUTED})],
               [("Aider: GPT-5.4-mini", {})], [("OpenHands: GPT-5.4-mini and Sonnet 4.5", {})],
               [("Codex: Sonnet 4.5", {})],
               [("Each can rerun the shown test. 40 minutes per run.", {"color": ON_DARK_MUTED})]])
text(s, XO + 0.2, CY + 0.52, WO - 0.4, 0.26, ["Repaired code under src/, judged on a clean copy"], color=MUTED)
rows = [["Setup", "P1", "P2", "P3 profit"],
        ["Aider + GPT-5.4-mini", "FAIL", "FAIL", "−0.098 ETH"],
        ["OpenHands + GPT-5.4-mini", "FAIL", "FAIL", "48.88 ETH"],
        ["OpenHands + Sonnet 4.5", "FAIL", "FAIL", "−0.098 ETH"],
        ["Codex + Sonnet 4.5", "FAIL", "FAIL", "−0.098 ETH"]]
colors = {(i, j): GREEN for i in range(1, 5) for j in (1, 2, 3)}
colors[(2, 3)] = RED
table(s, XO + 0.2, CY + 0.84, [2.02, 0.56, 0.56, 0.98], rows, row_h=0.34, size=11, colors=colors)
text(s, XO + 0.2, CY + 2.62, WO - 0.4, 0.95, [
    [("OpenHands + GPT-5.4-mini: ", {"bold": True, "color": INK}),
     ("\"Preserved exclusive behavior for same-word ranges\"", {})],
    [("Codex + Sonnet 4.5: ", {"bold": True, "color": INK}),
     ("\"my change is correct according to the requirement\"", {})]], color=MUTED, after=4)
band(s, "RESULT", "All 4 repairs make the shown test pass. 3 of 4 fix the bug. 1 of 4 hides it: P1 and P2 fail, "
     "but P3 still steals 48.88 ETH. All 18 hook tests pass on every repair. Cost: $9.62 for the 3 runs with a price.")

# ---------- step 4, actual data ----------
FILLED = "pilot/runs/openhands_openai/63008-V1-repair-P2/repair_prompt.md"
out = read(FILLED).split("Output on the current code:")[1].split("```text\n")[1].split("\n```")[0].split("\n")
fill = {"$requirement": "[... the V1 requirement, word for word ...]", "$test_path": "test/SwapFromTickBoundary.t.sol",
        "$test_code": "[... the test, on the right ...]", "$test_output": f"[... {len(out)} lines of output, on the right ...]",
        "$test_cmd": "forge test --match-path test/SwapFromTickBoundary.t.sol --match-test test_swapDownFromTickBoundary -vvv"}
lines = []
for l in prompt_text("pilot/prompts/repair_with_failing_test.md", "## Prompt"):
    if l.startswith("```"):
        continue  # the code fences around the test, the command and the output
    for k, v in fill.items():
        l = l.replace(k, v)
    lines.append(l)
s = new_slide("AI STEP 4 OF 5 · ACTUAL DATA · THE REPAIR PROMPT",
              "The repair prompt, with the failing test and its output",
              "Source: the repair prompt file (pilot/prompts/repair_with_failing_test.md) and the prompt sent to "
              "OpenHands + GPT-5.4-mini (its repair_prompt.md)",
              "This is the whole repair prompt. The pipeline fills in the requirement, the test, the command and the "
              "output of the test on the agent's own V1 code. The shown test is P2 with vm.expectRevert() and its "
              "comments removed, so it fails while the swap reverts. Every setup gets the same text. Only the "
              "output differs, because it comes from each agent's own code.")
wl = 5.9
verbatim(s, L, 1.86, wl, BOTTOM - 1.86, lines, size=8.5, gap=2)
fail_line = next(l for l in out if l.startswith("[FAIL"))
head, rest = fail_line.split(", ", 1)
yt = code_panel(s, L + wl + 0.15, 1.86, W - wl - 0.15, file_lines("pilot/tasks/63008/pov/P2_feedback_test.t.sol", 8, 24),
                title="The shown test, lines 8 to 24 of 24", size=8)
code_panel(s, L + wl + 0.15, yt + 0.14, W - wl - 0.15, [
    (" ", next(l for l in out if l.startswith("Ran 1 test"))), (" ", head + ","),
    (" ", "  " + rest.split(" (gas")[0]), (" ", "  (gas" + rest.split(" (gas")[1]), ("gap", "..."),
    (" ", next(l.strip() for l in out if "ArithmeticOverflowUnderflow" in l))],
    title=f"Its output on the agent's V1 code: 3 of {len(out)} lines", size=8)

RS = ("Source: each repair's diff against its own V1 code (pilot/runs/<setup>/63008-V1-repair-P2/repair_change.patch), "
      "cut to one line of context")
s, yb = diff_slide("AI STEP 4 OF 5 · ACTUAL DATA · THE SHALLOW REPAIR",
               "The shallow repair includes the start tick only in a narrow case", RS,
               "This diff compares the repaired code with the agent's own V1 code. The return inside the new if stops "
               "reset before the first step, so the first getNext returns the start tick. That happens only when all "
               "three conditions hold. In every other case the code falls through to its old walk, which skips the "
               "start tick. The agent kept running the human's iterator tests, which expect the start tick to be "
               "skipped, and its summary says \"the iterator tests still pass\".", [
    (REPAIR_PATCH % "openhands_openai", "OpenHands + GPT-5.4-mini", "BUG STAYS", False, [
        "The new return includes the start tick only if all three hold:",
        "1. the start and end ticks are in different bitmap words;",
        "2. the start tick is initialized and has net liquidity;",
        "3. the end tick is not a multiple of the tick spacing.",
        "Otherwise the old walk runs, and it skips the start tick."])])
band(s, "THE RESULT", "P1 and P2 fail, so the code looks fixed, but P3 still makes 48.88 ETH. Its summary names the "
     "gap without seeing it: \"Preserved exclusive behavior for same-word ranges\". P3's taxed swap starts and ends "
     "in the same bitmap word.", fill=RED_BG, line="F2B8B2", heading_color=RED, y=yb + 0.2)

s = new_slide("AI STEP 4 OF 5 · ACTUAL DATA · WHY P3 STILL WORKS",
              "P1 and P2 meet all three conditions. P3's swap does not",
              "Source: the three PoVs (pilot/tasks/63008/pov/) and the repaired code (pilot/runs/openhands_openai/"
              "63008-V1-repair-P2/repair_change.patch)",
              "The hook walks from the tick before the swap down to the tick after it. All three PoVs use tick "
              "spacing 10. A bitmap word holds 256 tick spacings, so ticks 0 to 2,559 are word 0, and ticks -2,560 "
              "to -1 are word -1. P1 and P2 swap from a positive tick to a negative one, so they cross from word 0 into "
              "word -1, and the repair includes the start tick. P3's attacker swaps from tick 20 to tick 19, inside "
              "word 0, so the repair falls back to the old walk and skips tick 20. The end ticks come from the price "
              "limit of each swap.")
text(s, L, 1.86, W, 0.5, ["The repaired reset includes the start tick only if all three conditions hold for the "
                          "swap that pays the tax. With tick spacing 10, one bitmap word covers ticks 0 to 2,559, and "
                          "the next word down covers ticks -2,560 to -1."], color=MUTED, size=11, line=14)
Y, N = GREEN, RED
table(s, L, 2.44, [0.6, 2.75, 1.9, 2.25, 1.45, 1.05, 1.55], [
    ["PoV", "The taxed swap: start tick to end tick", "1. Different words?", "2. Start tick has net liquidity?",
     "3. End off the spacing?", "Start tick", "On this repair"],
    ["P1", "30 to -51: down to the price of tick -51", "yes: word 0 to word -1", "yes: the top of [10, 30)",
     "yes: -51", "included", "FAIL: no revert"],
    ["P2", "20 to -11: down to the price of tick -10", "yes: word 0 to word -1", "yes: the top of [-10, 20)",
     "yes: -11", "included", "FAIL: no revert"],
    ["P3", "20 to 19: the attacker's swap back down", "no: both in word 0", "yes: the top of the attacker's [10, 20)",
     "yes: 19", "skipped", "PASS: 48.88 ETH profit"]],
    row_h=0.42, size=10.5, colors={(1, 5): Y, (1, 6): Y, (2, 5): Y, (2, 6): Y, (3, 2): N, (3, 5): N, (3, 6): N})
cw2 = (W - 0.15) / 2
card(s, L, 4.5, cw2, 1.72)
label(s, L + 0.2, 4.66, cw2 - 0.4, "WHY P1 AND P2 STOP WORKING", color=GREEN)
text(s, L + 0.2, 5.0, cw2 - 0.4, 1.6, [
    "In both tests, the start tick is the top of a position. If the walk skips it, the hook later subtracts "
    "liquidity it never added, the result goes below zero, and the swap reverts. That revert is what P1 and P2 "
    "check for. Here the start tick is included, so the swap goes through and both tests fail."], size=11, line=14)
card(s, L + cw2 + 0.15, 4.5, cw2, 1.72)
label(s, L + cw2 + 0.35, 4.66, cw2 - 0.4, "WHY P3 STILL PAYS THE ATTACKER", color=RED)
text(s, L + cw2 + 0.35, 5.0, cw2 - 0.4, 1.6, [
    "A wide position over [0, 50) holds the honest users' liquidity on both sides of tick 20. So nothing goes below "
    "zero, and nothing reverts. But the walk still skips tick 20, so the hook books the tax with the wrong active "
    "liquidity. The attacker's narrow position then takes the honest users' rewards: 48.88 ETH of the 49 ETH they "
    "built up."], size=11, line=14)
diff_slide("AI STEP 4 OF 5 · ACTUAL DATA · TWO REPAIRS THAT FIX THE BUG",
           "Two Sonnet 4.5 repairs let the first step return the start tick", RS,
           "Both diffs compare the repaired code with the agent's own V1 code. Both change where the walk starts, so "
           "the start tick comes back from the first getNext when it is initialized. P1, P2 and P3 fail on both, and "
           "P3 loses 0.098 ETH, the same as on the fix commit.", [
    (REPAIR_PATCH % "openhands_anthropic", "OpenHands + Sonnet 4.5", "FIXED", True, [
        "Returns before the first step when the start tick is initialized, so the first getNext returns it."]),
    (REPAIR_PATCH % "codex_anthropic", "Codex + Sonnet 4.5", "FIXED", True, [
        "Starts the search one tick spacing above the start tick's compressed position, so the first "
        "step down can return the start tick.",
        "It saw a human test fail and kept the fix: \"my change is correct according to the requirement, and it "
        "fixes the vulnerability\"."])])
diff_slide("AI STEP 4 OF 5 · ACTUAL DATA · THE AIDER REPAIR",
           "Aider + GPT-5.4-mini rewrote the start of the walk", RS,
           "This diff compares the repaired code with the agent's own V1 code. Aider applied no edit in its first "
           "three runs because of a code fence problem, which run_task.py now avoids. This is the fourth run, the "
           "counted one. Its first edit did not compile, and its second passed the shown test.", [
    (REPAIR_PATCH % "aider_openai", "Aider + GPT-5.4-mini", "FIXED", True, [
        "Finds the start tick itself when it is initialized, and walks down from there when it is not.",
        "P1, P2 and P3 fail. 10 of the 14 downward iterator tests pass. The 4 that fail expect the start tick to "
        "be skipped."])])

# ---------- step 5, judge on the repairs ----------
s = new_slide("AI STEP 5 OF 5 · LLM JUDGE ON THE REPAIRED CODE · 19 OCTOBER 2026",
              "After one check, the judge agreed with the PoVs on all 4 repairs",
              "Source: agent results, LLM judge (pilot/tasks/63008/feature_tasks/results.md), and the judge replies "
              "(pilot/runs/security_review/63008-V1-repair-P2/)",
              "The judge uses prompt 4 again, unchanged. Its known findings include finding 63008 itself, so the "
              "judge knows the bug and the agent did not. The audited code is reviewed too. Both judges find the "
              "bug there, so a 'no flaw' on a repair means something. GPT-5.4 made one wrong report with high "
              "confidence. A test of its own scenario rejected it.")
three_columns(s, "INPUT TO EACH JUDGE", "THE MODELS")
text(s, XI + 0.2, CY + 0.56, WI - 0.4, 2.0, [
    "Prompt 4 again, unchanged, with the V1 requirement",
    "The file with this version of reset(TickIteratorDown), and every caller",
    "The known findings, which include finding 63008 itself. So the judge knows the bug, and the agent did not."],
    bullets=True, after=5)
text(s, XI + 0.2, CY + 2.2, WI - 0.4, 1.1, [
    "5 versions: the audited code 386baff, which has the bug, and the 4 repairs. The audited code shows "
    "whether a judge can find the bug at all."], color=MUTED)
model_text(s, [[("Claude Opus 5.5", {"bold": True, "face": HEADF, "size": 14})],
               [("GPT-5.4", {"bold": True, "face": HEADF, "size": 14})],
               [("Same settings as AI step 3", {"color": ON_DARK_MUTED})],
               [("10 calls", {"color": ON_DARK_MUTED})],
               [("Opus: 13 to 23 seconds per review", {"color": ON_DARK_MUTED})],
               [("GPT-5.4: 138 to 164 seconds per review", {"color": ON_DARK_MUTED})]])
rows = [["Version", "Opus 5.5", "GPT-5.4"],
        ["Audited 386baff", "63008", "63008"],
        ["Aider + GPT-5.4-mini", "no flaw", "63008, rejected"],
        ["OpenHands + GPT-5.4-mini", "63008", "63008"],
        ["OpenHands + Sonnet 4.5", "no flaw", "no flaw"],
        ["Codex + Sonnet 4.5", "no flaw", "no flaw"]]
colors = {(1, 1): RED, (1, 2): RED, (3, 1): RED, (3, 2): RED, (2, 2): AMBER}
table(s, XO + 0.2, CY + 0.56, [1.98, 0.84, 1.3], rows, row_h=0.33, size=11, colors=colors)
quote(s, XO + 0.2, CY + 2.68, WO - 0.4, 0.82, [
    [("Opus on the shallow repair: ", {"bold": True}),
     ("\"Initialized start tick is skipped unless start and end are in different bitmap words and the end tick "
      "is unaligned\"", {})]], size=10.5)
band(s, "THE CHECK", "GPT-5.4's report on the Aider repair was turned into a test of its own "
     "scenario. The swap reverts on the audited code and goes through on the repair, so the report is rejected. "
     "After this check the judge agrees with the PoVs on all 4 repairs.", sub="by Claude Opus 5.5")

# ---------- step 5, actual data ----------
s = new_slide("AI STEP 5 OF 5 · ACTUAL DATA · PROMPT 4, REUSED",
              "The judge gets prompt 4 again, word for word",
              "Source: prompt 4 in the review prompt file (pilot/prompts/security_review.md). Blue text is filled in "
              "for each version.",
              "This is the same prompt as in AI step 3. Only the filled-in parts differ: V1's requirement, the repaired "
              "downward reset, and its callers. The known findings are the same five, so finding 63008 is in the "
              "prompt. The judge therefore knows the bug, which the agent did not.")
x3, w3 = prompt_columns(s, prompt_text("pilot/prompts/security_review.md", "## Prompt 4"), size=10, last_h=2.7)
card(s, x3, 4.7, w3, BOTTOM - 4.7, fill=AMBER_BG, line="F0D39A")
label(s, x3 + 0.18, 4.84, w3 - 0.36, "FILLED IN FOR THE JUDGE", color=AMBER)
text(s, x3 + 0.18, 5.16, w3 - 0.36, BOTTOM - 5.26, [
    [("<requirement>: ", {"face": MONO, "size": 9.5, "color": BLUE}), ("the V1 requirement", {})],
    [("<file>: ", {"face": MONO, "size": 9.5, "color": BLUE}), ("TickIterator.sol with the repaired reset", {})],
    [("<known_findings>: ", {"face": MONO, "size": 9.5, "color": BLUE}),
     ("the same five, so finding 63008 is in the prompt", {})],
    "The audited code is judged too, to check that a judge can find the bug at all."], size=10.5, after=3, line=13)
o = json.load(open(REVIEW % ("V1-repair-P2", "openhands_openai", OPUS)))
f = o["review"]["flaws"][0]
s = new_slide("AI STEP 5 OF 5 · ACTUAL DATA · THE JUDGE ON THE SHALLOW REPAIR",
              "Claude Opus 5.5 named the case where the repair still fails",
              "Source: the judge replies on the OpenHands + GPT-5.4-mini repair (pilot/runs/security_review/"
              "63008-V1-repair-P2/openhands_openai/)",
              f"This is the whole reply from Claude Opus 5.5 on the shallow repair. It took {o['seconds']} seconds. "
              "It reports finding 63008 and states the condition under which the start tick is still skipped. That "
              "matches what P3 shows. GPT-5.4 also reported finding 63008 on this repair.")
yc = code_panel(s, L, 1.86, W, [(" ", l) for l in f["lines"].split("\n")],
                title="\"lines\": the code it quotes from the repaired reset")
reply(s, L, yc + 0.14, 5.65, BOTTOM - yc - 0.14, [
    ("title", f["title"], False), ("severity", f["severity"], False), ("confidence", f["confidence"], False),
    ("known_finding", f["known_finding"], False), ("impact", f["impact"], False),
    ("requirement_breaks", o["review"]["requirement_breaks"], False)])
reply(s, L + 5.8, yc + 0.14, W - 5.8, BOTTOM - yc - 0.14, [("sequence", f["sequence"], False)])

g = json.load(open(REVIEW % ("V1-repair-P2", "aider_openai", GPT)))
f = g["review"]["flaws"][0]
CHECK = "pilot/tasks/63008/repair_checks/"
s = new_slide("AI STEP 5 OF 5 · ACTUAL DATA · THE REJECTED REPORT",
              "GPT-5.4's report on the Aider repair failed its own scenario",
              "Source: GPT-5.4's reply (pilot/runs/security_review/63008-V1-repair-P2/aider_openai/), the test and "
              "logs (pilot/tasks/63008/repair_checks/)",
              "GPT-5.4 reported finding 63008 in the Aider repair with high confidence. Step 3 of its sequence says "
              "reset \"can leave\" the iterator below tick 30, but it does not say which line does it. Claude Opus "
              "5.5 wrote a test of exactly the state in step 1. The swap reverts on the audited code and goes through "
              "on the repair, so the report is rejected. P1, P2 and P3 also fail on this repair.")
reply(s, L, 1.86, 5.5, BOTTOM - 1.86, [
    ("title", f["title"], False), ("severity", f["severity"], False), ("confidence", f["confidence"], False),
    ("known_finding", f["known_finding"], False),
    ("lines", f"[... the {len(f['lines'].splitlines())} lines of the repaired reset, left out here ...]", False),
    ("sequence", f["sequence"], False), ("impact", f["impact"], False)])
xr = L + 5.65
yt = code_panel(s, xr, 1.86, W - 5.65, file_lines(CHECK + "SwapDownFromTick30_GPT54Report.t.sol", 17, 29),
                title="The test of its scenario, written by Claude Opus 5.5", size=8)
fl = failed(CHECK + "swap_down_from_tick30_audited_386baff.log")
head, rest = fl.split(", ", 1)
passed = next(l.strip() for l in read(CHECK + "swap_down_from_tick30_aider_openai_repair.log").split("\n")
              if l.strip().startswith("[PASS"))
code_panel(s, xr, yt + 0.14, W - 5.65, [
    ("c", "# audited code 386baff, which has the bug"), (" ", head + ","), (" ", "  " + rest),
    ("c", "# the Aider + GPT-5.4-mini repair"), (" ", passed)],
    title="Results, in Docker with the network off", size=8)

# ---------- what checks each AI output ----------
s = new_slide("DEEP DIVE · WHAT CHECKS EACH AI OUTPUT",
              "Most AI output has a test. No person has read the reports yet",
              "Source: limits in the agent results (pilot/tasks/63008/feature_tasks/results.md) and the control "
              "group notes (pilot/tasks/63008/control_group.md)",
              "Every AI output except the requirements has a check that runs in Docker. But some of those checks "
              "were also written by Claude Opus 5.5, and a person has not yet read the review reports, the PoVs "
              "that confirm them, or the repaired code. Each cell is one run, so the numbers show what happened "
              "once, not a rate.")
rows = [["AI output", "Made by", "Checked by", "Still to do"],
        ["3 requirements", "Claude Opus 5.5", "Nothing yet. The drafter had read the finding.",
         "A blind run of prompt 2, then prompt 3"],
        ["Agent code for V1, 6 runs", "Coding agents", "P1, P2 and P3, in Docker",
         "A functional test. The project's tests expect the bug."],
        ["Agent code for B1 and B2, 12 runs", "Coding agents", "14 upward tests, 18 hook tests, blind review",
         "A person reads the code"],
        ["Review reports, 26 calls", "Claude Opus 5.5, GPT-5.4", "2 PoVs written by Claude Opus 5.5",
         "A person reads the reports and PoVs, and sets the severity"],
        ["Repaired code, 4 runs", "Coding agents", "P1, P2, P3 and the 18 hook tests",
         "A person reads the code against the requirement"],
        ["Judge reports, 10 calls", "Claude Opus 5.5, GPT-5.4", "PoVs, and 1 test written by Claude Opus 5.5",
         "A person reads the reports"]]
table(s, L, 1.92, [2.75, 2.15, 3.3, 3.35], rows, row_h=0.5)
band(s, "ONE MODEL", "Claude Opus 5.5 wrote the requirements, was one of the two reviewers, and "
     "wrote the tests that checked the reviews. Sonnet 4.5, one of the models under study, comes from the same "
     "company.", sub="in three roles", fill=AMBER_BG, line="F0D39A", heading_color=AMBER)

prs.save(DST)
print(f"wrote {DST}: {len(prs.slides)} slides")
